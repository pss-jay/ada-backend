"""PDF text extraction service for SOW documents."""

import logging
from pathlib import Path

logger = logging.getLogger(__name__)


def extract_text_from_pdf(filepath: str) -> str:
    """Extract all text from a PDF file using pdfplumber.

    Args:
        filepath: Path to the PDF file.

    Returns:
        The full extracted text content.
    """
    try:
        import pdfplumber

        text_parts = []
        with pdfplumber.open(filepath) as pdf:
            for page_num, page in enumerate(pdf.pages, 1):
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
                logger.debug(f"Extracted page {page_num}: {len(page_text or '')} chars")

        full_text = "\n\n".join(text_parts)
        logger.info(f"Extracted {len(full_text)} characters from {len(text_parts)} pages")
        return full_text

    except ImportError:
        logger.warning("pdfplumber not available, trying basic extraction")
        return _fallback_extract(filepath)
    except Exception as e:
        logger.error(f"Error extracting text from PDF: {e}", exc_info=True)
        raise


def extract_text_from_docx(filepath: str) -> str:
    """Extract text from a DOCX file.

    Args:
        filepath: Path to the DOCX file.

    Returns:
        The full extracted text content.
    """
    try:
        from docx import Document

        doc = Document(filepath)
        text_parts = []

        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)

        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    text_parts.append(row_text)

        full_text = "\n".join(text_parts)
        logger.info(f"Extracted {len(full_text)} characters from DOCX")
        return full_text

    except Exception as e:
        logger.error(f"Error extracting text from DOCX: {e}", exc_info=True)
        raise


def extract_text(filepath: str) -> str:
    """Extract text from a file (PDF or DOCX) based on extension."""
    path = Path(filepath)
    ext = path.suffix.lower()

    if ext == ".pdf":
        return extract_text_from_pdf(filepath)
    elif ext in (".docx", ".doc"):
        return extract_text_from_docx(filepath)
    else:
        raise ValueError(f"Unsupported file type: {ext}. Only PDF and DOCX are supported.")


def _fallback_extract(filepath: str) -> str:
    """Fallback text extraction when pdfplumber is not available."""
    # Try PyPDF2 as fallback
    try:
        from PyPDF2 import PdfReader

        reader = PdfReader(filepath)
        text_parts = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                text_parts.append(text)
        return "\n\n".join(text_parts)
    except ImportError:
        raise RuntimeError(
            "No PDF extraction library available. "
            "Install pdfplumber: pip install --break-system-packages pdfplumber"
        )
